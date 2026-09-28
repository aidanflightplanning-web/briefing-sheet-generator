"""Reader for the monthly MEL/CDL status workbooks (B737 and ATR).

Each workbook has one worksheet per day ("23RD SEPTEMBER 2026"). A sheet has a
header row (ACFT | MEL | CAT | ST | EXP | ... | DESCRIPTION | ... | CLOSURE # |
DATE) followed by one block of rows per aircraft: the three-letter registration
is in the first column (merged over the block) and each deferred item is a row
with an MEL number. Rectified items keep their row but get a closure number
and date and are coloured green ("GREEN: ENTIRE ROW OF CLEARED ITEMS"); those
are not carried onto the briefing sheet.
"""

from __future__ import annotations

import io
import re
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

MONTHS = {
    "JANUARY": 1, "FEBRUARY": 2, "MARCH": 3, "APRIL": 4, "MAY": 5, "JUNE": 6, "JULY": 7,
    "AUGUST": 8, "SEPTEMBER": 9, "OCTOBER": 10, "NOVEMBER": 11, "DECEMBER": 12,
    "JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "JUN": 6, "JUL": 7, "AUG": 8, "SEP": 9,
    "SEPT": 9, "OCT": 10, "NOV": 11, "DEC": 12,
}
SHEET_DATE_RE = re.compile(r"(\d{1,2})\s*(?:ST|ND|RD|TH)?\s+([A-Z]{3,9})\.?,?\s+(\d{4})")
REG_RE = re.compile(r"^(?:9Y-?)?([A-Z]{3})$")


@dataclass
class MelItem:
    registration: str
    number: str
    category: str
    start: date | None
    expiry: date | None
    description: str

    @property
    def briefing_text(self) -> str:
        return f"{self.number}  {self.description}"


@dataclass
class MelSheet:
    name: str
    date: date | None
    title: str
    aircraft: list[str] = field(default_factory=list)
    open_items: dict[str, list[MelItem]] = field(default_factory=dict)


@dataclass
class MelWorkbook:
    path: Path
    fleet: str
    sheets: list[MelSheet]

    def sheet_named(self, name: str) -> MelSheet | None:
        return next((s for s in self.sheets if s.name == name), None)

    def sheet_for(self, on: date | None) -> tuple[MelSheet | None, bool]:
        """Sheet for the given day; else the latest earlier sheet; else the last sheet.

        Returns (sheet, exact_match).
        """
        dated = [s for s in self.sheets if s.date]
        if on and dated:
            exact = [s for s in dated if s.date == on]
            if exact:
                return exact[-1], True
            earlier = [s for s in dated if s.date < on]
            if earlier:
                return max(earlier, key=lambda s: s.date), False
        if dated:
            return max(dated, key=lambda s: s.date), False
        return (self.sheets[-1] if self.sheets else None), False


@dataclass
class MelLookup:
    registration: str
    items: list[MelItem]
    workbook: MelWorkbook | None
    sheet: MelSheet | None
    exact_date: bool

    @property
    def found(self) -> bool:
        return self.sheet is not None

    @property
    def source(self) -> str:
        if not self.sheet:
            return "not found in the attached MEL sheets"
        return f"{self.workbook.path.name} > {self.sheet.name}"


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return re.sub(r"\s+", " ", str(value)).strip()


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _is_green(cell) -> bool:
    color = getattr(getattr(cell, "font", None), "color", None)
    rgb = getattr(color, "rgb", None)  # theme/indexed colours give a non-hex value
    if not isinstance(rgb, str) or not re.fullmatch(r"[0-9A-Fa-f]{6}(?:[0-9A-Fa-f]{2})?", rgb):
        return False
    r, g, b = (int(rgb[-6:][i:i + 2], 16) for i in (0, 2, 4))
    return g >= 0x80 and g - r >= 0x30 and g - b >= 0x30


def _sheet_date(ws) -> tuple[date | None, str]:
    title = ""
    for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
        for value in row:
            if isinstance(value, str) and "MEL" in value.upper():
                title = _text(value)
                break
        if title:
            break
    for text in (ws.title, title):
        m = SHEET_DATE_RE.search(text.upper())
        if m and m.group(2) in MONTHS:
            try:
                return date(int(m.group(3)), MONTHS[m.group(2)], int(m.group(1))), title
            except ValueError:
                pass
    return None, title


def _find_header(ws) -> tuple[int, dict[str, int]] | None:
    for r in range(1, min(ws.max_row, 15) + 1):
        labels = {c: _text(ws.cell(r, c).value).upper() for c in range(1, min(ws.max_column, 30) + 1)}
        mel_col = next((c for c, t in labels.items() if t in ("MEL", "MEL #", "MEL NO", "MEL/CDL")), None)
        desc_col = next((c for c, t in labels.items() if t.startswith("DESCRIPTION")), None)
        if not (mel_col and desc_col):
            continue
        cols = {"mel": mel_col, "desc": desc_col}
        cols["acft"] = next((c for c, t in labels.items() if t in ("ACFT", "A/C", "AIRCRAFT", "REG")), 1)
        cols["cat"] = next((c for c, t in labels.items() if t in ("CAT", "CAT.")), 0)
        cols["start"] = next((c for c, t in labels.items() if t in ("ST", "START", "DATE ST")), 0)
        cols["exp"] = next((c for c, t in labels.items() if t.startswith("EXP")), 0)
        closure = next((c for c, t in labels.items() if "CLOSURE" in t), 0)
        cols["closure"] = closure
        cols["closed_date"] = next(
            (c for c, t in labels.items() if closure and c > closure and "DATE" in t), 0)
        return r, cols
    return None


def _parse_sheet(ws) -> MelSheet:
    sheet_date, title = _sheet_date(ws)
    sheet = MelSheet(name=ws.title.strip(), date=sheet_date, title=title)
    header = _find_header(ws)
    if not header:
        return sheet
    header_row, cols = header

    merged_reg: dict[int, str] = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col <= cols["acft"] <= rng.max_col and rng.min_row > header_row:
            value = _text(ws.cell(rng.min_row, rng.min_col).value).upper().replace(" ", "")
            for r in range(rng.min_row, rng.max_row + 1):
                merged_reg[r] = value

    current = None
    for r in range(header_row + 1, ws.max_row + 1):
        row_text = [_text(ws.cell(r, c).value).upper() for c in range(1, min(ws.max_column, 30) + 1)]
        if "LEGEND" in row_text:
            break
        acft = merged_reg.get(r) or _text(ws.cell(r, cols["acft"]).value).upper().replace(" ", "")
        if acft:
            m = REG_RE.match(acft)
            current = m.group(1) if m else None
            if current and current not in sheet.aircraft:
                sheet.aircraft.append(current)
        if not current:
            continue
        mel_cell = ws.cell(r, cols["mel"])
        number = _text(mel_cell.value)
        if not number or number.upper() in ("N/A", "NIL", "-"):
            continue
        closure = _text(ws.cell(r, cols["closure"]).value) if cols["closure"] else ""
        closed_on = ws.cell(r, cols["closed_date"]).value if cols["closed_date"] else None
        if closure or closed_on or _is_green(mel_cell):
            continue
        item = MelItem(
            registration=current,
            number=number,
            category=_text(ws.cell(r, cols["cat"]).value) if cols["cat"] else "",
            start=_as_date(ws.cell(r, cols["start"]).value) if cols["start"] else None,
            expiry=_as_date(ws.cell(r, cols["exp"]).value) if cols["exp"] else None,
            description=_text(ws.cell(r, cols["desc"]).value).upper(),
        )
        sheet.open_items.setdefault(current, []).append(item)
    return sheet


_cache: dict[tuple[str, float], MelWorkbook] = {}


def load_workbook(source: str | Path | bytes, name: str | None = None) -> MelWorkbook:
    """Load and parse a MEL workbook from a file path (cached until the file changes)
    or from the workbook's bytes (not cached; `name` names it)."""
    import openpyxl

    if isinstance(source, (bytes, bytearray)):
        path, key, stream = Path(name or "MEL sheet.xlsx"), None, io.BytesIO(source)
    else:
        path = Path(source)
        key, stream = (str(path.resolve()), path.stat().st_mtime), path
        if key in _cache:
            return _cache[key]
    with warnings.catch_warnings():  # openpyxl warns about Excel extensions it does not need
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(stream, data_only=True)
    try:
        sheets = [_parse_sheet(ws) for ws in wb.worksheets]
    finally:
        wb.close()
    name = path.name.upper()
    titles = " ".join(s.title.upper() for s in sheets[:1])
    fleet = "ATR" if "ATR" in name or "ATR" in titles else "B737" if "737" in name or "737" in titles else ""
    workbook = MelWorkbook(path=path, fleet=fleet, sheets=sheets)
    if key:
        _cache[key] = workbook
    return workbook


def lookup(workbooks: list[MelWorkbook], registration: str, on: date | None,
           sheet_name: str | None = None) -> MelLookup:
    """Open MEL items for one aircraft, read from the sheet for the briefing day."""
    reg = re.sub(r"^9Y-?", "", registration.upper())
    for wb in workbooks:
        if not any(reg in s.aircraft for s in wb.sheets):
            continue
        sheet, exact = (wb.sheet_named(sheet_name), True) if sheet_name else wb.sheet_for(on)
        if sheet is None:
            sheet, exact = wb.sheet_for(on)
        if sheet is None:
            continue
        return MelLookup(reg, list(sheet.open_items.get(reg, [])), wb, sheet, exact)
    return MelLookup(reg, [], None, None, False)
