"""Tests with synthetic data: a small flight plan brief and MEL workbook are built in memory.

Run from the project folder:  py -m unittest -v
Checks against real example flight plans live in test_local_examples.py (kept out of the repository).
"""

import io
import unittest
from datetime import date, datetime

from briefsheet import flightplan, mel, render
from briefsheet.builder import BriefingData, build_briefing
from briefsheet.notams import Notam, flight_remarks
from briefsheet.output import output_names, package_pdf
from briefsheet.pdftext import extract_pages
from briefsheet.taf import format_taf


def make_pdf(pages: list[list], header: bool = True) -> bytes:
    """A PDF laid out like the brief. A line is plain text (next row, Courier 9) or
    (text, font size, x, row) to place it exactly, as the NOTAM table needs."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    for number, lines in enumerate(pages, start=1):
        row = 0
        for line in lines:
            text, size, x, row = (line, 9, 40, row) if isinstance(line, str) else line
            c.setFont("Courier-Bold" if size > 10 else "Courier", size)
            c.drawString(x, 750 - 11 * row, text)
            row += isinstance(line, str)
        if header:
            c.setFont("Courier", 8)
            c.drawString(360, 780, f"CARIBBEAN AIRLINES LIMITED BRIEF PAGE {number} OF {len(pages)}")
            c.drawString(520, 20, f"PAGE {number} OF {len(pages)}")
        c.showPage()
    c.save()
    return buffer.getvalue()


NOTAM_PAGE_TOP = [("TTPP (PIARCO INTL) - TBPB (GRANTLEY ADAMS INTL)", 13, 154, 0),
                  ("Briefing generated: 01OCT26 0230 (All times in UTC unless otherwise stated) FLT BWA0123", 8, 32, 2),
                  ("NTM No: Eff from:Eff to: Description:", 8, 32, 3)]


def airport(row: int, name: str) -> list[tuple]:
    return [(name, 11, 32, row)]


def category(row: int, name: str) -> list[tuple]:
    return [(name, 11, 175, row)]


def notam(row: int, number: str, start, end, *text: str) -> list[tuple]:
    """The cells of one NOTAM in the order the brief draws them: number, from, to, description."""
    cells = [(number, 9.5, 32, row)]
    if start == "WIE":
        cells.append(("WIE", 9.5, 87, row))
    else:
        cells += [(start[0], 9.5, 87, row), (start[1], 9.5, 87, row + 1)]
    if isinstance(end, str):
        cells.append((end, 9.5, 131, row))
    else:
        cells += [(end[0], 9.5, 131, row), (end[1], 9.5, 131, row + 1)]
    return cells + [(line, 9.5, 175, row + i) for i, line in enumerate(text)]


BRIEF = [
    ["RMKS", "RAIM CHECK VALIDATED", "XFUE DUE EN RTE WX",
     "FLTNO /DATE ORIG/DEST ALTN EQUIP/REG PYLD FUEL RAMP B/O LDWT",
     "BWA123 /01OCT POS /BGI GND B737 /ABC 0100 0070 0600 0020 0580",
     "SCHED DEP 1200/1245 ARR UTC ARAMPWT .......",
     "TOTAL FUEL REQUIRED I 02.10 006123",
     "DIVERSION TO GND", "ALTN1 TIME 00.30 FL100 B/O 000400 CMR 0800KGS DIST 0100 WCP P005 TDV P15",
     "DIVERSION TO TAB", "ALTN2 TIME 00.20 FL080 B/O 000300 CMR 0700KGS DIST 0060 WCP P005 TDV P15",
     "EQUAL TIME POINT / ENGINEOUT DEPRESSURIZATION TIME . . . FOB . . .",
     "TTPP TO TGPY POSN N11 00.0 W061 00.0 GWT 060000KGS DDALT 10000FT",
     "TGPY TO TBPB POSN N12 00.0 W060 00.0 GWT 059000KGS DDALT 10000FT",
     "(FPL-BWA123-IS", "-B38M/M-SDE3FGHIRWYZ/LB1", "-TTPP1200", "-N0450F280 DCT POS UR515 BGI DCT",
     "-TBPB0040 TGPY TTCP", "-PBN/A1B1C1D1S2 DOF/261001 REG/9YABC)", "DISPATCHER - JANE DOE"],
    ["DEPARTURE", "METAR TTPP 011100Z 09010KT 9999 FEW016 28/24 Q1012",
     "TAF TTPP 011000Z 0112/0212 09010KT 9999 SCT018 PROB30 TEMPO", "0114/0118 4000 SHRA BKN014",
     "ARRIVAL", "TAF TBPB 011000Z 0112/0212 10012KT 9999 SCT020",
     "OTHER", "TAF TGPY 011000Z 0112/0212 10008KT 9999 FEW018 BECMG 0200/0202 VRB03KT",
     "OTHER", "TAF TTCP 011000Z 0112/0212 09010KT 9999 FEW018"],
    NOTAM_PAGE_TOP
    + airport(5, "PIARCO INTL (TTPP/POS), TRINIDAD TOBAGO PIARCO FIR (TTZP)") + [("ELEV. 58FT", 9.5, 175, 6)]
    + category(8, "COMPANY NOTAMS") + notam(9, "00001", "WIE", "UFN", "COMPANY FREQ : 123.45")
    + category(11, "RUNWAY")
    + notam(12, "A0001/26", ("01OCT26", "1100"), ("01OCT26", "1400"), "RWY 10/28 CLSD DUE WIP")
    + notam(15, "A0002/26", ("01OCT26", "2000"), ("02OCT26", "0600"), "RWY 10/28 CLSD")
    + category(18, "TAXIWAY(S)")
    + notam(19, "A0003/26", ("01OCT26", "0000"), ("03OCT26", "0000"), "TWY A CLSD BTN RWY 10 AND TWY B")
    + airport(22, "GRANTLEY ADAMS INTL (TBPB/BGI), BARBADOS") + [("ELEV. 169FT", 9.5, 175, 23)]
    + category(25, "LOCALIZER (ILS)")
    + notam(26, "A0010/26", ("30SEP26", "0000"), "UFN", "ILS LOC 'IBGI' 110.100MHZ RWY 09"),   # text goes on overleaf
    NOTAM_PAGE_TOP
    + airport(5, "GRANTLEY ADAMS INTL (TBPB/BGI), BARBADOS") + [("U/S DUE MAINT", 9.5, 175, 6)]
    + category(8, "RUNWAY")
    + notam(9, "A0011/26", ("28SEP26", "0300"), ("05OCT26", "1000"), "0300 - 1000", "RWY 09/27 CLSD")
    + category(12, "PRECISION APPROACH PATH INDICATOR")
    + notam(13, "A0012/26", ("01OCT26", "0000"), ("03OCT26", "0000"), "RWY 09 PAPI U/S")
    + airport(16, "ST. GEORGES/MAURICE BISHOP - I (TGPY/GND), GRENADA") + category(17, "AERODROME")
    + notam(18, "A0020/26", ("01OCT26", "1400"), ("01OCT26", "1800"), "AD CLSD DUE VIP MOVEMENT")
    + airport(21, "SCARBOROUGH/A.N.R.ROBINSON - I (TTCP/TAB), TRINIDAD TOBAGO")
    + category(22, "INSTRUMENT APPROACH PROCEDURE")
    + notam(23, "A0030/26", ("01OCT26", "0000"), ("03OCT26", "0000"), "RNAV (GNSS) RWY 11 APCH NOT AVBL")
    + airport(26, "MIAMI INTL (KMIA/MIA), UNITED STATES") + category(27, "RUNWAY")
    + notam(28, "A0040/26", ("01OCT26", "0000"), ("03OCT26", "0000"), "RWY 09/27 CLSD")
    + airport(31, "PIARCO INTL (TTPP/POS), TRINIDAD TOBAGO") + [("SEE ABOVE", 9.5, 175, 32)],
    ["TTZP SIGMET 1 VALID 010100/010500 TTPP-", "TTZP PIARCO FIR EMBD TS OBS AT 0050Z WI N1100 W06230",
     "TOP FL400 MOV W 10KT NC"],
]
NOTAM_REMARKS = [
    "TTPP 01OCT26 1100-1400Z: RWY 10/28 CLSD DUE WIP [A0001/26]",
    "TBPB 30SEP26 0000Z-UFN: ILS LOC 'IBGI' 110.100MHZ RWY 09 U/S DUE MAINT [A0010/26]",
    "TGPY 01OCT26 1400-1800Z: AD CLSD DUE VIP MOVEMENT [A0020/26]",
    "TTCP 01OCT26 0000Z-03OCT26 0000Z: RNAV (GNSS) RWY 11 APCH NOT AVBL [A0030/26]",
]


class TafLayout(unittest.TestCase):
    def test_change_groups(self):
        raw = ("TAF TTPP 232200Z 2400/2424 09003KT 9999 FEW010CB SCT015 TEMPO 2402/2409 5000 SHRA BKN015 "
               "BECMG 2411/2413 11010KT PROB30 TEMPO 2415/2418 3000 TSRA SCT010CB")
        self.assertEqual(format_taf(raw), [
            "TAF TTPP 232200Z 2400/2424 09003KT 9999 FEW010CB SCT015",
            "    TEMPO 2402/2409 5000 SHRA BKN015",
            "  BECMG 2411/2413 11010KT",
            "  PROB30",
            "    TEMPO 2415/2418 3000 TSRA SCT010CB",
        ])

    def test_us_style_from_and_prob_with_time(self):
        raw = ("TAF KIAD 260541Z 2606/2712 32009G16KT P6SM SCT250 FM261400 35016G26KT P6SM SCT130 BKN250 "
               "PROB30 2709/2712 5SM -RA BKN035")
        self.assertEqual(format_taf(raw), [
            "TAF KIAD 260541Z 2606/2712 32009G16KT P6SM SCT250",
            "  FM261400 35016G26KT P6SM SCT130 BKN250",
            "  PROB30 2709/2712 5SM -RA BKN035",
        ])

    def test_remarks_stay_inline(self):
        raw = "TAF CYYZ 232340Z 2400/2506 09012KT P6SM FEW260 BECMG 2422/2424 VRB03KT RMK NXT FCST BY 240300Z"
        self.assertEqual(format_taf(raw)[-1], "  BECMG 2422/2424 VRB03KT RMK NXT FCST BY 240300Z")


class ParserEdgeCases(unittest.TestCase):
    def test_sigmet_directly_after_the_weather(self):
        lines = ["DEPARTURE", "METAR TTPP 240200Z 10004KT 9999 FEW016 28/25 Q1012 NOSIG",
                 "TAF TTPP 232200Z 2400/2424 09003KT 9999 FEW010CB", "SCT015",
                 "TJZS SIGMET GOLF 2 VALID 240205/240605 KKCI-", "SAN JUAN OCEANIC FIR FRQ TS"]
        tafs, _ = flightplan._parse_weather(lines)
        self.assertEqual(tafs["TTPP"], "TAF TTPP 232200Z 2400/2424 09003KT 9999 FEW010CB SCT015")
        self.assertEqual(flightplan._parse_sigmets(lines),
                         [["TJZS SIGMET GOLF 2 VALID 240205/240605 KKCI-", "SAN JUAN OCEANIC FIR FRQ TS"]])

    def test_theme_font_colour_is_not_treated_as_green(self):
        class Color:
            rgb = "Values must be of type <class 'str'>"

        class Font:
            color = Color()

        class Cell:
            font = Font()

        self.assertFalse(mel._is_green(Cell()))

    def test_atr_flight_line(self):
        m = flightplan.FLTNO_RE.match("BWA1540 /30SEP POS /TAB POS ATR72/TTA 6200 1370 0211 0256 20799")
        self.assertEqual((m.group("flight"), m.group("orig"), m.group("dest"), m.group("altn"), m.group("equip"),
                          m.group("reg")), ("BWA1540", "POS", "TAB", "POS", "ATR72", "TTA"))
        m = flightplan.FLTNO_RE.match("BWA5 /26SEP KIN /JFK IAD B737 /BAR 0170 0130 0757 0081 0676")
        self.assertEqual((m.group("number"), m.group("equip"), m.group("reg")), ("5", "B737", "BAR"))


class SyntheticBrief(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pdf = make_pdf(BRIEF)
        cls.fp = flightplan.parse_flight_plan(cls.pdf, name="BWA012301TTPP.pdf")
        cls.data = build_briefing(cls.fp, gate="5", pax=["3/120"])

    def test_leg(self):
        (leg,) = self.fp.legs
        self.assertEqual((leg.flight, leg.orig, leg.dest, leg.registration), ("BWA123", "POS", "BGI", "ABC"))
        self.assertEqual((leg.dep_icao, leg.dest_icao, leg.alternates), ("TTPP", "TBPB", ["TGPY", "TTCP"]))
        self.assertEqual(leg.total_fuel, 6123)
        self.assertEqual((leg.departure, leg.arrival), (datetime(2026, 10, 1, 12, 0), datetime(2026, 10, 1, 12, 45)))
        self.assertEqual(leg.alternate_minutes, [30, 20])
        self.assertEqual(self.fp.path.name, "BWA012301TTPP.pdf")
        self.assertEqual(self.fp.briefing_generated, datetime(2026, 10, 1, 2, 30))

    def test_briefing_fields(self):
        d = self.data
        self.assertEqual((d.aircraft_reg, d.flight_no, d.gate, d.pax_load), ("ABC", "123", "5", "3/120"))
        self.assertEqual(d.flight_plan_fuel, "6200KG")
        self.assertEqual(d.alt_filed, "TGPY & TTCP")
        self.assertEqual(d.wx_departure.splitlines(), [
            "TAF TTPP 011000Z 0112/0212 09010KT 9999 SCT018", "  PROB30", "    TEMPO 0114/0118 4000 SHRA BKN014"])
        self.assertEqual(d.wx_destination, "TAF TBPB 011000Z 0112/0212 10012KT 9999 SCT020")
        self.assertEqual(d.wx_alternate.splitlines()[0], "TAF TGPY 011000Z 0112/0212 10008KT 9999 FEW018")
        self.assertIn("TAF TTCP", d.wx_alternate)
        self.assertEqual(d.sigmet_info.splitlines()[0], "TTZP SIGMET 1 VALID 010100/010500 TTPP-")
        self.assertEqual(d.sigmet_info.splitlines()[-1], " TOP FL400 MOV W 10KT NC")
        self.assertEqual(d.enroute_airport, "TGPY")
        self.assertEqual(d.dispatcher, "JANE DOE")
        self.assertEqual(d.date, "30/09/26")   # 0230Z on 1 Oct is still 30 Sep in Trinidad

    def test_notams_are_read_per_airport(self):
        notams = self.fp.notams
        self.assertEqual({k: [n.number for n in v] for k, v in notams.items()}, {
            "TTPP": ["00001", "A0001/26", "A0002/26", "A0003/26"], "TBPB": ["A0010/26", "A0011/26", "A0012/26"],
            "TGPY": ["A0020/26"], "TTCP": ["A0030/26"], "KMIA": ["A0040/26"]})
        company, closed = notams["TTPP"][0], notams["TTPP"][1]
        self.assertEqual((company.start, company.end, company.open_end, company.category, company.text),
                         (None, None, "UFN", "COMPANY NOTAMS", "COMPANY FREQ : 123.45"))
        self.assertEqual((closed.start, closed.end, closed.category),
                         (datetime(2026, 10, 1, 11, 0), datetime(2026, 10, 1, 14, 0), "RUNWAY"))
        ils = notams["TBPB"][0]      # until further notice, and its text continues on the next page
        self.assertEqual((ils.end, ils.open_end, ils.category), (None, "UFN", "LOCALIZER (ILS)"))
        self.assertEqual(ils.text, "ILS LOC 'IBGI' 110.100MHZ RWY 09 U/S DUE MAINT")
        nightly = notams["TBPB"][1]
        self.assertEqual((nightly.daily, nightly.text), ([(180, 600)], "RWY 09/27 CLSD"))

    def test_closures_and_outages_for_the_flight_go_into_the_remarks(self):
        # Selected: the departure runway closure (inside dep-1h .. arr+1h), the destination ILS, the
        # alternate's closure (starts within its 30 min diversion time) and the second alternate's approach.
        # Not selected: a closure that evening, a taxiway, a nightly closure outside the flight time,
        # a PAPI, and a runway closed at an airport the flight does not use.
        self.assertEqual(flight_remarks(self.fp), NOTAM_REMARKS)
        self.assertEqual(self.data.additional_info.splitlines(),
                         ["• XFUE DUE EN RTE WX"] + ["• " + remark for remark in NOTAM_REMARKS])

    def test_warns_when_notams_cannot_be_checked(self):
        self.assertEqual(self.fp.warnings, [])
        no_times = [[line for line in BRIEF[0] if not line.startswith("SCHED DEP")], *BRIEF[1:]]
        fp = flightplan.parse_flight_plan(make_pdf(no_times))
        self.assertEqual(fp.warnings, ["BWA123 POS-BGI: scheduled times not found - NOTAMs were not checked for this leg."])
        self.assertEqual(flight_remarks(fp), [])
        # Without the larger heading font the airports cannot be told apart from the NOTAM text.
        plain = ["TTPP (PIARCO INTL) - TBPB (GRANTLEY ADAMS INTL)", "Briefing generated: 01OCT26 0230 (UTC) FLT BWA0123",
                 "NTM No: Eff from:Eff to: Description:", "PIARCO INTL (TTPP/POS), TRINIDAD TOBAGO", "RUNWAY",
                 "A0001/26 01OCT26", "1100", "01OCT26", "1400", "RWY 10/28 CLSD"]
        fp = flightplan.parse_flight_plan(make_pdf([BRIEF[0], BRIEF[1], plain]))
        self.assertEqual(fp.notams, {})
        self.assertEqual(fp.warnings, ["The NOTAM pages could not be read - closures and outages were not added "
                                       "to the remarks."])

    def test_package_replaces_an_attached_sheet(self):
        old_sheet = make_pdf([["Flight Crew Briefing Sheet", "FORM NBR SOC001"]], header=False)
        from pypdf import PdfReader, PdfWriter

        writer = PdfWriter()
        for source in (self.pdf, old_sheet):
            for page in PdfReader(io.BytesIO(source)).pages:
                writer.add_page(page)
        with_old = io.BytesIO()
        writer.write(with_old)
        new_sheet = io.BytesIO()
        render.render_briefing(self.data, new_sheet)
        texts = extract_pages(package_pdf(with_old.getvalue(), new_sheet.getvalue()))
        self.assertEqual(len(texts), len(BRIEF) + 1)
        self.assertIn("GATE POSITION", texts[-1])
        self.assertIn("RWY 10/28 CLSD DUE WIP", texts[-1])
        self.assertEqual(output_names("BWA012301TTPP_WITH_BRIEFSHEET.pdf"),
                         ("BWA012301TTPP_BRIEFSHEET.pdf", "BWA012301TTPP_WITH_BRIEFSHEET.pdf"))


class NotamScreening(unittest.TestCase):
    """Wording taken from real NOTAMs: what is a closure or outage, and what is not."""

    def kind(self, category, text):
        return Notam("TTPP", "A0001/26", category=category, lines=[text]).kind

    def test_closures_and_outages(self):
        for category, text, kind in [
            ("RUNWAY", "RWY 06L/24R CLSD AVBL AS TWY.", "RWY CLSD"),
            ("RUNWAY", "TUE-SUN 0200-1000 RWY 10/28 CLSD NIGHTLY TO ALL OPS EXC EMERG. 90 MIN PPR.", "RWY CLSD"),
            ("UNCATEGORISED", "JFK RWY 13R/31L CLSD", "RWY CLSD"),
            ("AERODROME", "AD AP CLSD TO NON SKED ACFT WINGSPAN MORE THAN 214FT", "AD CLSD"),
            ("AERODROME CONTROL TOWER (TWR)", "V.C. BIRD AD CONTROL TOWER CLSD DUE STAFF SHORTAGE.", "TWR CLSD"),
            ("COMPANY NOTAMS", "TTPP A1037/23 15JUL0027-13OCT2359 NDB 'TRI' 382.0KHZ U/S", "NAVAID"),
            ("UNCATEGORISED", "SJU NAV VORTAC U/S", "NAVAID"),
            ("VOR", "HAMPTON VOR 'HTO' FREQ 113.6MHZ U/S.", "NAVAID"),
            ("GLIDE PATH (ILS)", "PHL NAV ILS RWY 27R GP U/S", "NAVAID"),
            ("LOCATOR OUTER (ILS)", "SJU NAV ILS RWY 08 PATTY LOM U/S", "NAVAID"),
            ("ILS", "NAV ILS RWY 27 UNUSABLE DA TO TDZ", "NAVAID"),
            ("LOCALIZER (ILS)", "ILS LOC 'IBGI' 110.100MHZ RWY 09 SUBJ INTRP DUE GRASS CUTTING. SGL MAY BE UNREL. "
                                "DO NOT USE.", "NAVAID"),
            ("INSTRUMENT APPROACH PROCEDURE", "IAP JOHN F KENNEDY INTL, NEW YORK, NY. VOR RWY 4R, ORIG-C... PROC NA.",
             "APCH"),
            ("INSTRUMENT APPROACH PROCEDURE", "DUE RWY CONST: RNAV (GNSS) Z RWY 06 AND RNAV (GNSS) Z RWY 24 APCH: "
                                              "NOT AUTH EXC FOR TRAINING FLT.", "APCH"),
        ]:
            self.assertEqual(self.kind(category, text), kind, text)

    def test_everything_else_is_left_to_the_notam_pages(self):
        for category, text in [
            ("TAXIWAY(S)", "PHL TWY S BTN TWY D AND APCH END RWY 27R CLSD"),
            ("UNCATEGORISED", "PHL TWY S BTN TWY D AND APCH END RWY 27R CLSD"),
            ("TAXIWAY(S)", "SECN OF TWY ALPHA (WEST OF TWY DELTA TO RWY 07 STARTER EXTENSION) CLSD DUE WORKS."),
            ("APRON", "APN HLDG PAD FOR RWY 09R CLSD"),
            ("RUNWAY", "RWY 10/28 WIP. MEN/EQPT OPR ON A PULL BACK BASIS."),
            ("RUNWAY", "RESA RWY 25 UNAVBL DUE CONST WORKS ON TWY A"),
            ("RUNWAY", "CONCENTRATION OF BIRDS IN THE VCY OF RWY 11/29. PLEASE BE ADVISED."),
            ("RUNWAY CENTRELINE LIGHTS", "RCLL RWY 09/27 U/S."),
            ("PRECISION APPROACH PATH INDICATOR", "RWY 08 PAPI U/S"),
            ("APPROACH LIGHT SYSTEM", "ALS 06R AND 24L U/S. APPLY PROC FOR HIGH INTST APCH LGT U/S (AIP AD 2.22.4)"),
            ("DAYLIGHT MARKINGS", "PHL TWY E HLDG PSN SIGN FOR APCH END RWY 35 LGT U/S"),
            ("UNCATEGORISED", "TWY E ELEVATED RWY GUARD LGT AT APCH END RWY 35 U/S"),
            ("APPROACH CONTROL SERVICE (APP)", "APP FREQUENCIES 119.000MHZ AND 119.550MHZ U/S. FREQ 124.000MHZ TO "
                                               "BE USED FOR APP"),
            ("INNER MARKER (ILS)", "NAV ILS RWY 04R IM U/S"),
            ("ILS", "IAD ILS RWY 19C IM U/S"),
            ("ILS", "NAV ILS RWY 27R SPECIAL AUTH CAT II NA"),
            ("INSTRUMENT APPROACH PROCEDURE", "IAP JOHN F KENNEDY INTL, NEW YORK, NY. ILS OR LOC RWY 4L, AMDT 11E... "
                                              "CHANGE NOTE: AUTOPILOT COUPLED APCH NA BLW 200 FT MSL"),
            ("INSTRUMENT APPROACH PROCEDURE", "SJU IAP SAN JUAN, PR. ILS OR LOC RWY 8, AMDT 16A ... ALTERNATE MISSED "
                                              "APCH NA"),
            ("INSTRUMENT APPROACH PROCEDURE", "IAP SAN JUAN, PR. VOR OR TACAN RWY 8, AMDT 1A ... DME REQUIRED EXC FOR "
                                              "ACFT EQUIPPED WITH SUITABLE RNAV SYSTEM WITH GPS, BQN VORTAC OUT OF SVC"),
            ("UNCATEGORISED", "SJU NAV VORTAC NOT MNT"),
            ("UNCATEGORISED", "IAD COM REMOTE TRANS/REC 128.525 U/S USE 127.325"),
            ("UNCATEGORISED", "AD AP WINDCONE FOR RWY 27R LGT U/S"),
            ("FIRE FIGHTING AND RESCUE", "AD AP ARFF VEHICLE U/S INDEX UNCHANGED"),
            ("COMPANY NOTAMS", "COMPANY FREQ : 123.45 COMPANY CONTACT NUMBER : 555-0100"),
        ]:
            self.assertIsNone(self.kind(category, text), text)

    def test_time_window_and_daily_schedule(self):
        nightly = Notam("MYNN", "A0001/26", datetime(2026, 9, 28, 22, 0), datetime(2026, 10, 5, 10, 0),
                        category="RUNWAY", lines=["2200-1000", "RWY 14/32 CLSD"])
        self.assertEqual(nightly.timeframe, "28SEP26 2200Z-05OCT26 1000Z DLY 2200-1000")
        self.assertTrue(nightly.active(datetime(2026, 10, 1, 8, 0), datetime(2026, 10, 1, 9, 30)))     # before 1000
        self.assertTrue(nightly.active(datetime(2026, 10, 1, 21, 0), datetime(2026, 10, 1, 22, 30)))   # reaches 2200
        self.assertFalse(nightly.active(datetime(2026, 10, 1, 11, 0), datetime(2026, 10, 1, 20, 0)))   # daytime
        self.assertFalse(nightly.active(datetime(2026, 10, 6, 1, 0), datetime(2026, 10, 6, 3, 0)))     # expired
        weekly = Notam("MYNN", "A0002/26", datetime(2026, 9, 17, 13, 28), datetime(2026, 9, 28, 10, 0),
                       category="RUNWAY", lines=["TUE-SUN 0200-1000 RWY 10/28 CLSD NIGHTLY"])
        self.assertIsNone(weekly.daily)         # not a plain daily schedule: treated as in force throughout
        self.assertTrue(weekly.active(datetime(2026, 9, 24, 12, 0), datetime(2026, 9, 24, 15, 0)))
        open_ended = Notam("KEWR", "A4747/25", datetime(2025, 10, 24, 5, 56), datetime(2026, 10, 21, 20, 0),
                           lines=["/COL/ DME U/S"])
        self.assertEqual(open_ended.timeframe, "24OCT25 0556Z-21OCT26 2000Z")
        self.assertEqual(Notam("TTPP", "00004", open_end="UFN", lines=["NDB U/S"]).remark, "TTPP WIE-UFN: NDB U/S [00004]")


class SyntheticMel(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import openpyxl
        from openpyxl.styles import Font

        book = openpyxl.Workbook()
        headers = ["", "MEL", "CAT", "ST", "EXP", "LP/SNG#/DWC", "AUTH", "DESCRIPTION", "REP. MAINT ACTION",
                   "PARTS", "REMARKS", "TROUBLESHOOTING", "DEFECT RECTIFICATION", "OPS", "CLOSURE #", "DATE"]
        for ws, day, open_item in ((book.active, "1ST OCTOBER 2026", ("21-00-01", "PACK TEMP LIGHT INOP")),
                                   (book.create_sheet(), "30TH SEPTEMBER 2026", ("30-11-01", "WING ANTI-ICE INOP"))):
            ws.title = day
            ws["A1"] = f"B737 MEL STAT/CDL STATUS  {day} UPDATE AT 1500 HRS"
            for col, text in enumerate(headers, start=1):
                ws.cell(4, col, text)
            ws["A5"], ws["B5"], ws["C5"], ws["H5"] = "ABC", open_item[0], "C", open_item[1]
            ws["B6"], ws["H6"], ws["O6"], ws["P6"] = "25-00-01", "TRAY TABLE 3A", "C01/XX/261001", datetime(2026, 10, 1)
            ws["B7"], ws["H7"] = "33-11-01", "LOGO LIGHT INOP"
            ws["B7"].font = Font(color="FF00B050")          # cleared: green row, closure # missing
            ws.merge_cells("A5:A7")
            ws["A8"], ws["G9"], ws["H9"] = "DEF", "NO", "ROBBERIES"
            ws["G11"] = "LEGEND"
            ws["A12"], ws["B12"], ws["H12"] = "ABC", "99-99-99", "NOT AN ITEM"
        buffer = io.BytesIO()
        book.save(buffer)
        cls.books = [mel.load_workbook(buffer.getvalue(), name="OCTOBER 2026 B737 MEL SHEET.xlsx")]

    def items(self, reg, day):
        return [i.briefing_text for i in mel.lookup(self.books, reg, day).items]

    def test_only_open_items_of_the_day(self):
        self.assertEqual(self.books[0].fleet, "B737")
        self.assertEqual(self.items("ABC", date(2026, 10, 1)), ["21-00-01  PACK TEMP LIGHT INOP"])
        self.assertEqual(self.items("ABC", date(2026, 9, 30)), ["30-11-01  WING ANTI-ICE INOP"])
        self.assertEqual(self.items("DEF", date(2026, 10, 1)), [])

    def test_missing_day_and_unknown_aircraft(self):
        later = mel.lookup(self.books, "ABC", date(2026, 10, 5))
        self.assertEqual((later.sheet.name, later.exact_date), ("1ST OCTOBER 2026", False))
        self.assertFalse(mel.lookup(self.books, "XYZ", date(2026, 10, 1)).found)
        chosen = mel.lookup(self.books, "ABC", date(2026, 10, 1), "30TH SEPTEMBER 2026")
        self.assertEqual(chosen.sheet.name, "30TH SEPTEMBER 2026")


def template_data() -> BriefingData:
    """The values on the Word template (Briefsheet.pdf)."""
    return BriefingData(
        aircraft_reg="GUY", flight_no="476-477", gate="13", pax_load="2/11 / 15/144", mel_items="NIL REPORTED",
        wx_departure="TAF TTPP 260400Z 2606/2706 00000KT 9999 FEW016\n  BECMG 2612/2614 11008KT 9999 SCT018\n"
                     "  PROB40\n    TEMPO 2614/2620 11010G20KT 3000 TSRA FEW010CB BKN015\n"
                     "  BECMG 2700/2702 00000KT 9999 FEW016 SCT038",
        wx_destination="TAF AMD MUHA 260854Z 2608/2706 VRB05KT 9000 FEW030\n    TEMPO 2609/2611 3000 BR\n"
                       "  BECMG 2614/2616 36008KT",
        wx_alternate="TAF MKJP 260515Z 2606/2706 VRB02KT 9999 FEW022\n  BECMG 2614/2616 16009KT\n"
                     "  BECMG 2623/2701 34004KT AMD NOT SKED\n\nTAF TBPB 260500Z 2606/2706 09014KT 9999 SCT018\n"
                     "  PROB30\n    TEMPO 2606/2612 SHRA BKN014\n\nTAF TGPY 260400Z 2606/2706 09008KT 9999 SCT018",
        flight_plan_fuel="19600KG / 11500KG", alt_filed="MKJP / TBPB & TGPY", dispatcher="A DISPATCHER",
        date="26/09/26")


class Rendering(unittest.TestCase):
    def words(self, data):
        import pdfplumber

        buffer = io.BytesIO()
        pages = render.render_briefing(data, buffer)
        with pdfplumber.open(io.BytesIO(buffer.getvalue())) as pdf:
            return pages, {w["text"]: w for w in pdf.pages[0].extract_words()}

    def test_layout_matches_template_positions(self):
        pages, words = self.words(template_data())
        self.assertEqual(pages, 1)
        # (text, x0, top) measured on the Word template
        for text, x0, top in [("GUY", 169.13, 184.03), ("476-477", 169.13, 204.79), ("13", 169.13, 223.87),
                              ("DOCUMENT", 77.66, 140.06), ("GATE", 77.66, 219.79), ("TBA", 169.13, 284.26),
                              ("STANDARD", 169.13, 306.70), ("DEPARTURE", 77.66, 337.18),
                              ("SIGMET", 77.66, 522.72), ("INFORMATION", 77.66, 606.15)]:
            self.assertAlmostEqual(words[text]["x0"], x0, delta=0.3, msg=text)
            self.assertAlmostEqual(words[text]["top"], top, delta=0.3, msg=text)

    def test_remarks_in_point_form(self):
        data = template_data()
        data.additional_info = "• TTPP 01OCT26 1100-1400Z: RWY 10/28 CLSD DUE WIP [A0001/26]"
        _, words = self.words(data)                       # a single NOTAM is still a bullet point
        self.assertAlmostEqual(words["TTPP"]["x0"], 169.13 + 36, delta=0.3)
        data.additional_info = "XFUE DUE EN RTE WX"
        _, words = self.words(data)                       # a single plain remark is not
        self.assertAlmostEqual(words["XFUE"]["x0"], 169.13, delta=0.3)

    def test_long_content_flows_onto_a_second_page(self):
        data = template_data()
        data.sigmet_info = "\n".join(["TJZS SIGMET GOLF 2 VALID 240205/240605 KKCI-"] + [" LINE"] * 20)
        self.assertEqual(render.render_briefing(data, io.BytesIO()), 2)

    def test_renders_without_arial_or_calibri_installed(self):
        saved = render._fonts, render._font_dirs, render._font_trees
        try:
            render._fonts = None
            render._font_dirs = render._font_trees = lambda: []
            self.assertEqual(render.render_briefing(template_data(), io.BytesIO()), 1)
            self.assertEqual(render.get_fonts(), render.Fonts("Helvetica", "Helvetica-Bold", "Helvetica-Bold"))
        finally:
            render._fonts, render._font_dirs, render._font_trees = saved


if __name__ == "__main__":
    unittest.main()
